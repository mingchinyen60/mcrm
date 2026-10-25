import io
import csv
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from typing import List, Dict, Any
import openpyxl
from app.core.database import get_db
from app.models.account import Account
from app.models.contact import Contact
from app.models.tech_spec import TechnicalSpec

router = APIRouter(prefix="/imports", tags=["Batch Uploads & Templates"])

# Template Definitions
TEMPLATES = {
    "accounts": {
        "headers": ["name", "type", "industry", "tier", "region", "website", "phone", "email", "address", "notes"],
        "sample": ["Arrow Electronics", "Distributor", "Global Distribution", "Strategic", "North America", "https://arrow.com", "+1 800 777 2776", "sales@arrow.com", "Centennial, CO", "Authorized global distribution partner"]
    },
    "contacts": {
        "headers": ["account_name", "first_name", "last_name", "title", "role_category", "email", "phone", "is_primary", "notes"],
        "sample": ["Arrow Electronics", "John", "Doe", "Supplier Business Manager", "Procurement", "jdoe@arrow.com", "+1 303 555 1234", "true", "Key disti line card contact"]
    },
    "tech_specs": {
        "headers": ["part_number", "category", "nominal_frequency", "frequency_stability", "package_type", "supply_voltage", "output_logic", "jitter_spec", "phase_noise_10hz", "phase_noise_100hz", "phase_noise_1khz", "phase_noise_10khz", "phase_noise_100khz", "phase_noise_1mhz", "target_application", "description"],
        "sample": ["OCXO-20M-HC40", "OCXO", "20.000 MHz", "±0.2 ppb", "HC40", "3.3V", "Sine Wave", "< 40 fs RMS", "-115", "-142", "-160", "-166", "-170", "-172", "IEEE 1588 PTP", "Low phase noise reference clock"]
    }
}

@router.get("/template/{entity_type}")
def download_template(entity_type: str, format: str = Query("csv", pattern="^(csv|excel)$")):
    if entity_type not in TEMPLATES:
        raise HTTPException(status_code=400, detail="Invalid template type. Use accounts, contacts, or tech_specs.")
        
    tpl = TEMPLATES[entity_type]
    headers = tpl["headers"]
    sample_row = tpl["sample"]

    if format == "excel":
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = f"{entity_type.capitalize()} Template"
        ws.append(headers)
        ws.append(sample_row)

        output = io.BytesIO()
        wb.save(output)
        output.seek(0)
        filename = f"{entity_type}_template.xlsx"
        return StreamingResponse(
            output,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )
    else:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(headers)
        writer.writerow(sample_row)
        output.seek(0)
        filename = f"{entity_type}_template.csv"
        return StreamingResponse(
            io.BytesIO(output.getvalue().encode("utf-8-sig")),
            media_type="text/csv",
            headers={"Content-Disposition": f"attachment; filename={filename}"}
        )

def parse_rows(file: UploadFile) -> List[Dict[str, Any]]:
    filename = file.filename.lower()
    content = file.file.read()

    rows = []
    if filename.endswith(".xlsx"):
        wb = openpyxl.load_workbook(io.BytesIO(content))
        ws = wb.active
        data = list(ws.iter_rows(values_only=True))
        if not data:
            return []
        headers = [str(h).strip().lower() if h is not None else "" for h in data[0]]
        for row in data[1:]:
            if not any(row):
                continue
            rows.append({headers[i]: row[i] for i in range(len(headers)) if i < len(row)})
    else:
        text = content.decode("utf-8-sig", errors="replace")
        reader = csv.DictReader(io.StringIO(text))
        for row in reader:
            clean_row = {k.strip().lower(): (v.strip() if isinstance(v, str) else v) for k, v in row.items() if k}
            rows.append(clean_row)
    return rows

@router.post("/accounts")
def import_accounts(file: UploadFile = File(...), db: Session = Depends(get_db)):
    rows = parse_rows(file)
    if not rows:
        raise HTTPException(status_code=400, detail="Uploaded file is empty or could not be parsed.")

    created = 0
    errors = []
    for idx, row in enumerate(rows, start=2):
        name = row.get("name")
        if not name:
            errors.append(f"Row {idx}: Missing account name")
            continue

        existing = db.query(Account).filter(Account.name.ilike(name)).first()
        if existing:
            # Update existing
            for k in ["type", "industry", "tier", "region", "website", "phone", "email", "address", "notes"]:
                if row.get(k):
                    setattr(existing, k, str(row[k]))
        else:
            acc = Account(
                name=str(name),
                type=str(row.get("type") or "OEM"),
                industry=str(row.get("industry") or "Telecommunications"),
                tier=str(row.get("tier") or "Tier 2"),
                region=str(row.get("region") or "North America"),
                website=str(row.get("website", "")) if row.get("website") else None,
                phone=str(row.get("phone", "")) if row.get("phone") else None,
                email=str(row.get("email", "")) if row.get("email") else None,
                address=str(row.get("address", "")) if row.get("address") else None,
                notes=str(row.get("notes", "")) if row.get("notes") else None,
            )
            db.add(acc)
            created += 1

    db.commit()
    return {"success": True, "created_count": created, "total_rows": len(rows), "errors": errors}

@router.post("/contacts")
def import_contacts(file: UploadFile = File(...), db: Session = Depends(get_db)):
    rows = parse_rows(file)
    if not rows:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    created = 0
    errors = []
    for idx, row in enumerate(rows, start=2):
        acc_name = row.get("account_name")
        first_name = row.get("first_name")
        last_name = row.get("last_name")

        if not acc_name or not first_name or not last_name:
            errors.append(f"Row {idx}: Missing account_name, first_name, or last_name")
            continue

        account = db.query(Account).filter(Account.name.ilike(acc_name)).first()
        if not account:
            # Create the account automatically if missing
            account = Account(name=acc_name)
            db.add(account)
            db.commit()
            db.refresh(account)

        is_prim = str(row.get("is_primary") or "false").lower() in ["true", "1", "yes"]
        c = Contact(
            account_id=account.id,
            first_name=str(first_name),
            last_name=str(last_name),
            title=str(row.get("title", "")) if row.get("title") else None,
            role_category=str(row.get("role_category") or "Engineering Lead"),
            email=str(row.get("email", "")) if row.get("email") else None,
            phone=str(row.get("phone", "")) if row.get("phone") else None,
            is_primary=is_prim,
            notes=str(row.get("notes", "")) if row.get("notes") else None
        )
        db.add(c)
        created += 1

    db.commit()
    return {"success": True, "created_count": created, "total_rows": len(rows), "errors": errors}

@router.post("/tech-specs")
def import_tech_specs(file: UploadFile = File(...), db: Session = Depends(get_db)):
    rows = parse_rows(file)
    if not rows:
        raise HTTPException(status_code=400, detail="Uploaded file is empty.")

    created = 0
    errors = []
    for idx, row in enumerate(rows, start=2):
        part_number = row.get("part_number")
        nominal_freq = row.get("nominal_frequency")
        stability = row.get("frequency_stability")

        if not part_number or not nominal_freq or not stability:
            errors.append(f"Row {idx}: Missing part_number, nominal_frequency, or frequency_stability")
            continue

        def to_float(val):
            if val is None:
                return None
            if isinstance(val, (int, float)):
                return float(val)
            try:
                s = str(val).strip()
                return float(s) if s != "" else None
            except:
                return None

        spec = db.query(TechnicalSpec).filter(TechnicalSpec.part_number == part_number).first()
        if not spec:
            spec = TechnicalSpec(
                part_number=str(part_number),
                category=str(row.get("category") or "OCXO"),
                nominal_frequency=str(nominal_freq),
                frequency_stability=str(stability),
                package_type=str(row.get("package_type") or "SMD 7050"),
                supply_voltage=str(row.get("supply_voltage") or "3.3V"),
                output_logic=str(row.get("output_logic") or "CMOS"),
                jitter_spec=str(row.get("jitter_spec", "")) if row.get("jitter_spec") else None,
                phase_noise_10hz=to_float(row.get("phase_noise_10hz")),
                phase_noise_100hz=to_float(row.get("phase_noise_100hz")),
                phase_noise_1khz=to_float(row.get("phase_noise_1khz")),
                phase_noise_10khz=to_float(row.get("phase_noise_10khz")),
                phase_noise_100khz=to_float(row.get("phase_noise_100khz")),
                phase_noise_1mhz=to_float(row.get("phase_noise_1mhz")),
                target_application=str(row.get("target_application") or "5G Infrastructure"),
                description=str(row.get("description", "")) if row.get("description") else None,
            )
            db.add(spec)
            created += 1
        else:
            # Update fields
            spec.nominal_frequency = nominal_freq
            spec.frequency_stability = stability
            if row.get("package_type"): spec.package_type = str(row["package_type"])
            if row.get("target_application"): spec.target_application = str(row["target_application"])

    db.commit()
    return {"success": True, "created_count": created, "total_rows": len(rows), "errors": errors}
